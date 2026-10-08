"""Shared builder for the OshiFit decks.

Builds on the lab template (Media Interaction Laboratory, Osaka Institute of
Technology): every deck starts from the template's master/layouts, the
template slides are dropped, and slides are composed from the helpers below.
"""
import os
import re
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.environ.get("OSHI_TEMPLATE", os.path.join(HERE, "template.pptx"))
ASSETS = os.path.join(HERE, "assets")

C = dict(
    navy="0B3C6E", blue="0099FF", sky="E6F4FF", sky2="CCE9FF", pink="E8457C", pinkL="FDE7EF",
    A="E07A10", AL="FFF1E3", B="12A38A", BL="E3F6F2", text="24303F", muted="5A6676",
    gray="7A8594", line="D5DEE8", white="FFFFFF", panel="F4F8FC", purple="7C6BC4", purpleL="EFEBFA",
    gold="F5B300", red="D64545",
)

L, R, TOP, BOT = 0.45, 9.55, 1.5, 6.85
W = R - L


def rgb(c):
    return RGBColor.from_string(C.get(c, c))


def asset(name):
    p = os.path.join(ASSETS, name)
    return p if os.path.exists(p) else os.path.join(ASSETS, name + ".png")


# ---------------------------------------------------------------- text
_MARK = re.compile(r"(\*\*.+?\*\*|\[\[.+?\]\]|\{\{.+?\}\})")


def _runs_from_markup(s):
    """'**bold**' -> bold, '[[x]]' -> bold pink, '{{x}}' -> bold navy."""
    out = []
    for part in _MARK.split(s):
        if not part:
            continue
        if part.startswith("**"):
            out.append((part[2:-2], dict(bold=True)))
        elif part.startswith("[["):
            out.append((part[2:-2], dict(bold=True, color="pink")))
        elif part.startswith("{{"):
            out.append((part[2:-2], dict(bold=True, color="navy")))
        else:
            out.append((part, {}))
    return out


def _set_bullet(p, char="•", color=None, indent=0.2, level=0):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent * (level + 1)))))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buClr", "a:buFont", "a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if color:
        bc = etree.SubElement(pPr, qn("a:buClr"))
        sc = etree.SubElement(bc, qn("a:srgbClr"))
        sc.set("val", C.get(color, color))
    bf = etree.SubElement(pPr, qn("a:buFont"))
    bf.set("typeface", "Arial")
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    if pPr.find(qn("a:buNone")) is None:
        etree.SubElement(pPr, qn("a:buNone"))


def fill_tf(tf, content, size=14, color="text", bold=False, align="l", anchor="t",
            margin=(0.05, 0.03, 0.05, 0.03), spacing=None, after=4, wrap=True, autofit=False):
    """content: str (\n separates paragraphs) or list of paragraph specs.

    A paragraph spec is a str (markup allowed), or a dict with keys
    t (str markup) | runs [(text, opts)], bullet (bool|char), level, size,
    color, bold, align, after, before, italic.
    """
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    if autofit:
        bp = tf._txBody.find(qn("a:bodyPr"))
        etree.SubElement(bp, qn("a:normAutofit"))
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margin]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(content, str):
        content = content.split("\n")
    first = True
    for spec in content:
        if isinstance(spec, str):
            spec = dict(t=spec)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        al = spec.get("align", align)
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[al]
        if spacing:
            p.line_spacing = spacing
        p.space_after = Pt(spec.get("after", after))
        if "before" in spec:
            p.space_before = Pt(spec["before"])
        b = spec.get("bullet")
        if b:
            _set_bullet(p, b if isinstance(b, str) else "•", spec.get("bcolor", "blue"),
                        indent=spec.get("indent", 0.2), level=spec.get("level", 0))
        else:
            _no_bullet(p)
        runs = spec.get("runs") or _runs_from_markup(spec.get("t", ""))
        sz = spec.get("size", size)
        for text, o in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.size = Pt(o.get("size", sz))
            f.bold = o.get("bold", spec.get("bold", bold))
            if o.get("italic", spec.get("italic")):
                f.italic = True
            f.color.rgb = rgb(o.get("color", spec.get("color", color)))
    return tf


# ---------------------------------------------------------------- deck
class Deck:
    def __init__(self, date, title="", author="今井元"):
        self.prs = Presentation(TEMPLATE)
        lst = self.prs.slides._sldIdLst
        for sid in list(lst):
            self.prs.part.drop_rel(sid.rId)
            lst.remove(sid)
        self.date = date
        self.prs.core_properties.title = title
        self.prs.core_properties.author = author
        self.layouts = {l.name.replace("\n", ""): l for l in self.prs.slide_layouts}
        self.n = 0

    # -- slides
    def title_slide(self, label, title, name, affil, subtitle=None, art=None, title_size=32):
        s = self.prs.slides.add_slide(self.layouts["タイトル スライド"])
        self.n += 1
        for ph in s.placeholders:
            idx = ph.placeholder_format.idx
            if idx == 0:
                fill_tf(ph.text_frame, [dict(t=t, align="c", after=0) for t in title.split("\n")], size=title_size,
                        color="text", anchor="b", margin=(0.1, 0.05, 0.1, 0.05))
            elif idx == 1:
                fill_tf(ph.text_frame, name, size=18, color="text", margin=(0.1, 0.05, 0.1, 0.05))
        # conference / meeting label on the blue band (as in the lab decks)
        lab = s.shapes.add_textbox(Inches(0.07), Inches(0.09), Inches(4.6), Inches(0.42))
        lab.fill.solid()
        lab.fill.fore_color.rgb = rgb("A9DCFF")
        fill_tf(lab.text_frame, label, size=14, color="text", anchor="m", margin=(0.1, 0.02, 0.1, 0.02))
        lab.name = "会議名"
        aff = s.shapes.add_textbox(Inches(4.3), Inches(5.86), Inches(5.4), Inches(1.2))
        fill_tf(aff.text_frame, affil, size=14, color="text", after=2)
        aff.name = "著者所属"
        if subtitle:
            st = s.shapes.add_textbox(Inches(0.45), Inches(1.0), Inches(6.2 if art else 9.1), Inches(2.3))
            fill_tf(st.text_frame, subtitle, size=20, color="white", bold=True, anchor="m", after=6, spacing=1.05)
            st.name = "サブタイトル"
        if art:
            self.pic(s, art, 6.85, 0.72, h=3.0)
        return s

    def slide(self, title, size=26):
        s = self.prs.slides.add_slide(self.layouts["タイトルのみ"])
        self.n += 1
        fill_tf(s.shapes.title.text_frame, title, size=size, color="text", anchor="m",
                margin=(0.2, 0.03, 0.1, 0.12))
        lay = s.slide_layout
        for ph in lay.placeholders:
            idx = ph.placeholder_format.idx
            if idx == 10:
                d = s.shapes.clone_placeholder(ph)
                d = s.placeholders[10]
                p = d.text_frame.paragraphs[0]
                r = p.add_run()
                r.text = self.date
            elif idx == 12:
                s.shapes.clone_placeholder(ph)
                sn = s.placeholders[12]
                p = sn.text_frame.paragraphs[0]._p
                fld = etree.SubElement(p, qn("a:fld"))
                fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
                fld.set("type", "slidenum")
                rpr = etree.SubElement(fld, qn("a:rPr"))
                rpr.set("lang", "en-US")
                t = etree.SubElement(fld, qn("a:t"))
                t.text = str(self.n)
        return s

    def save(self, path):
        self.prs.save(path)

    # -- primitives
    def text(self, s, x, y, w, h, content, name=None, fill=None, line=None, **kw):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        fill_tf(tb.text_frame, content, **kw)
        if fill:
            tb.fill.solid()
            tb.fill.fore_color.rgb = rgb(fill)
        if line:
            tb.line.color.rgb = rgb(line)
        if name:
            tb.name = name
        return tb

    def box(self, s, x, y, w, h, fill="panel", line=None, lw=1.0, radius=0.08, shape=None,
            text=None, name=None, shadow=False, **kw):
        if shape is None:
            shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE and radius:
            sh.adjustments[0] = min(0.5, radius / min(w, h))
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(lw)
        else:
            sh.line.fill.background()
        if shadow:
            self._soft_shadow(sh)
        else:
            sh.shadow.inherit = False
        if text is not None:
            kw.setdefault("color", "text")
            kw.setdefault("anchor", "m")
            fill_tf(sh.text_frame, text, **kw)
        else:
            sh.text_frame.text = ""
        if name:
            sh.name = name
        return sh

    def _soft_shadow(self, sh):
        spPr = sh._element.spPr
        eff = etree.SubElement(spPr, qn("a:effectLst"))
        sd = etree.SubElement(eff, qn("a:outerShdw"))
        sd.set("blurRad", "76200")
        sd.set("dist", "19050")
        sd.set("dir", "5400000")
        sd.set("algn", "t")
        sd.set("rotWithShape", "0")
        c = etree.SubElement(sd, qn("a:srgbClr"))
        c.set("val", "0B3C6E")
        a = etree.SubElement(c, qn("a:alpha"))
        a.set("val", "18000")

    def circle(self, s, x, y, d, fill, line=None, text=None, **kw):
        return self.box(s, x, y, d, d, fill=fill, line=line, shape=MSO_SHAPE.OVAL, radius=None, text=text, **kw)

    def icon(self, s, name, x, y, d=0.5, bg="blue", fg="w", pad=0.24):
        if bg:
            self.circle(s, x, y, d, bg)
        k = d * (1 - 2 * pad)
        return self.pic(s, f"ic_{name}_{fg}", x + (d - k) / 2, y + (d - k) / 2, w=k, h=k)

    def pic(self, s, name, x, y, w=None, h=None):
        return s.shapes.add_picture(asset(name), Inches(x), Inches(y),
                                    Inches(w) if w else None, Inches(h) if h else None)

    def line(self, s, x1, y1, x2, y2, color="gray", width=1.5, arrow=False, dash=False, head=False):
        cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        cn.line.color.rgb = rgb(color)
        cn.line.width = Pt(width)
        ln = cn.line._get_or_add_ln()
        if dash:
            pd = etree.SubElement(ln, qn("a:prstDash"))
            pd.set("val", "dash")
        if head:
            he = etree.SubElement(ln, qn("a:headEnd"))
            he.set("type", "triangle")
        if arrow:
            te = etree.SubElement(ln, qn("a:tailEnd"))
            te.set("type", "triangle")
            te.set("w", "med")
            te.set("len", "med")
        return cn

    def arrow(self, s, x, y, w, h, fill="blue", direction="r"):
        shp = {"r": MSO_SHAPE.RIGHT_ARROW, "d": MSO_SHAPE.DOWN_ARROW, "l": MSO_SHAPE.LEFT_ARROW,
               "u": MSO_SHAPE.UP_ARROW, "chev": MSO_SHAPE.CHEVRON}[direction]
        return self.box(s, x, y, w, h, fill=fill, shape=shp, radius=None)

    def pill(self, s, x, y, w, h, text, fill="blue", color="white", size=11, bold=True, line=None):
        return self.box(s, x, y, w, h, fill=fill, line=line, radius=h / 2, text=text, size=size, bold=bold,
                        color=color, align="c", margin=(0.04, 0.0, 0.04, 0.0))

    def assumed(self, s, text="想定値：実データで差し替え"):
        """Marks a slide whose numbers are projections, sitting on the footer band."""
        return self.pill(s, 5.9, 7.09, 3.05, 0.3, "※ " + text, fill="white", color="pink", size=10, line="pink")

    def notes(self, s, text):
        s.notes_slide.notes_text_frame.text = text

    def card(self, s, x, y, w, h, title, body, icon=None, color="blue", fill="white", tsize=15, bsize=12.5,
             line="line", shadow=True, title_color=None, body_after=3):
        self.box(s, x, y, w, h, fill=fill, line=line, radius=0.1, shadow=shadow)
        tx = x + 0.18
        if icon:
            self.icon(s, icon, x + 0.16, y + 0.16, 0.5, bg=color)
            tx = x + 0.78
            self.text(s, tx, y + 0.13, w - (tx - x) - 0.12, 0.56, title, size=tsize, bold=True,
                      color=title_color or "navy", anchor="m", margin=(0, 0, 0, 0))
            by = y + 0.76
        else:
            self.text(s, tx, y + 0.12, w - 0.3, 0.42, title, size=tsize, bold=True, color=title_color or color,
                      anchor="m", margin=(0, 0, 0, 0))
            by = y + 0.56
        self.text(s, x + 0.18, by, w - 0.32, h - (by - y) - 0.1, body, size=bsize, color="text", margin=(0, 0, 0, 0),
                  after=body_after)

    # -- table
    def table(self, s, x, y, w, rows, colw=None, rowh=0.36, size=12, header_fill="navy", header_color="white",
              zebra="panel", first_col_bold=False, align=None, hsize=None, bold_rows=(), fills=None, valign="m"):
        nr, nc = len(rows), len(rows[0])
        gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rowh * nr))
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is not None:
            sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid
        tbl.first_row = False
        tbl.horz_banding = False
        colw = colw or [w / nc] * nc
        for i, cw in enumerate(colw):
            tbl.columns[i].width = Inches(cw)
        for r in range(nr):
            tbl.rows[r].height = Inches(rowh)
            for c in range(nc):
                cell = tbl.cell(r, c)
                v = rows[r][c]
                cell.margin_left = cell.margin_right = Inches(0.07)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE if valign == "m" else MSO_ANCHOR.TOP
                hdr = r == 0 and header_fill
                al = (align[c] if align else "l")
                if hdr:
                    al = "c" if (align and align[c] == "c") else al
                spec = dict(t=str(v), align=al)
                fill_tf(cell.text_frame, [spec], size=(hsize or size) if hdr else size,
                        color=header_color if hdr else "text",
                        bold=bool(hdr) or (first_col_bold and c == 0) or (r in bold_rows),
                        margin=(0.07, 0.03, 0.07, 0.03), after=0, anchor=valign)
                cell.fill.solid()
                if hdr:
                    cell.fill.fore_color.rgb = rgb(header_fill)
                elif fills and fills.get((r, c)):
                    cell.fill.fore_color.rgb = rgb(fills[(r, c)])
                elif zebra and r % 2 == 0:
                    cell.fill.fore_color.rgb = rgb(zebra)
                else:
                    cell.fill.fore_color.rgb = rgb("white")
                self._cell_border(cell)
        return gf

    def _cell_border(self, cell, color="D5DEE8"):
        tcPr = cell._tc.get_or_add_tcPr()
        for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            ln = tcPr.find(qn(tag))
            if ln is None:
                ln = etree.Element(qn(tag))
                # borders must precede the fill in tcPr
                tcPr.insert(0, ln)
            ln.set("w", "9525")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            c = etree.SubElement(sf, qn("a:srgbClr"))
            c.set("val", color)
        # keep schema order: lnL, lnR, lnT, lnB, ..., solidFill
        order = [qn("a:lnL"), qn("a:lnR"), qn("a:lnT"), qn("a:lnB")]
        lns = [tcPr.find(t) for t in order]
        for el in lns:
            tcPr.remove(el)
        for i, el in enumerate(lns):
            tcPr.insert(i, el)

    # -- charts
    def _style_chart(self, ch, legend=True, size=11, title=None):
        ch.font.size = Pt(size)
        ch.font.color.rgb = rgb("muted")
        if title:
            ch.has_title = True
            ch.chart_title.text_frame.text = title
            tp = ch.chart_title.text_frame.paragraphs[0]
            tp.runs[0].font.size = Pt(13)
            tp.runs[0].font.bold = True
            tp.runs[0].font.color.rgb = rgb("navy")
        else:
            ch.has_title = False
        ch.has_legend = legend
        if legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(size)

    def _axes(self, ch, vmin=None, vmax=None, major=None, fmt=None, gridlines=True, val_visible=True):
        va = ch.value_axis
        if vmin is not None:
            va.minimum_scale = vmin
        if vmax is not None:
            va.maximum_scale = vmax
        if major:
            va.major_unit = major
        va.has_major_gridlines = gridlines
        if gridlines:
            va.major_gridlines.format.line.color.rgb = rgb("E3E8EF")
            va.major_gridlines.format.line.width = Pt(0.75)
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(10.5)
        va.tick_labels.font.color.rgb = rgb("muted")
        if fmt:
            va.tick_labels.number_format = fmt
            va.tick_labels.number_format_is_linked = False
        if not val_visible:
            va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        try:
            ca = ch.category_axis
            ca.format.line.color.rgb = rgb("B8C2CE")
            ca.tick_labels.font.size = Pt(11)
            ca.tick_labels.font.color.rgb = rgb("text")
            ca.has_major_gridlines = False
        except Exception:
            pass

    def bar(self, s, x, y, w, h, cats, series, colors, horizontal=False, stacked=False, labels=True, fmt="0.0",
            vmin=None, vmax=None, major=None, legend=None, title=None, gap=60, overlap=None, point_colors=None,
            label_pos="out", val_visible=True, size=11, label_color=None, cat_low=False):
        cd = CategoryChartData()
        cd.categories = cats
        for name, vals in series:
            cd.add_series(name, vals)
        if horizontal:
            t = XL_CHART_TYPE.BAR_STACKED if stacked else XL_CHART_TYPE.BAR_CLUSTERED
        else:
            t = XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED
        gf = s.shapes.add_chart(t, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        self._style_chart(ch, legend=(len(series) > 1) if legend is None else legend, size=size, title=title)
        self._axes(ch, vmin, vmax, major, fmt=fmt, val_visible=val_visible)
        plot = ch.plots[0]
        plot.gap_width = gap
        if overlap is not None:
            plot.overlap = overlap
        elif stacked:
            plot.overlap = 100
        for i, ser in enumerate(plot.series):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
            ser.invert_if_negative = False
        if point_colors:
            for i, pc in enumerate(point_colors):
                pt = plot.series[0].points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(pc)
        if labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.number_format = fmt
            dl.number_format_is_linked = False
            dl.font.size = Pt(size)
            dl.font.bold = True
            dl.font.color.rgb = rgb(label_color or ("white" if stacked or label_pos == "in" else "text"))
            dl.position = (XL_LABEL_POSITION.CENTER if stacked else
                           XL_LABEL_POSITION.INSIDE_END if label_pos == "in" else XL_LABEL_POSITION.OUTSIDE_END)
        if cat_low:
            ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        if horizontal:
            ch.category_axis.reverse_order = True
            from pptx.enum.chart import XL_TICK_LABEL_POSITION as TLP
            ch.value_axis.tick_label_position = TLP.NONE if not val_visible else TLP.HIGH
            # keep value labels at the bottom when the category axis is reversed
            crosses = ch.value_axis._element
        return ch

    def linechart(self, s, x, y, w, h, cats, series, colors, fmt="0.0", vmin=None, vmax=None, major=None,
                  labels=False, title=None, legend=True, widths=None, dashes=None, size=11, label_last=False):
        cd = CategoryChartData()
        cd.categories = cats
        for name, vals in series:
            cd.add_series(name, vals)
        gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        self._style_chart(ch, legend=legend, size=size, title=title)
        self._axes(ch, vmin, vmax, major, fmt=fmt)
        for i, ser in enumerate(ch.plots[0].series):
            col = colors[i % len(colors)]
            ser.smooth = False
            ser.format.line.color.rgb = rgb(col)
            ser.format.line.width = Pt((widths or [2.5] * 10)[i])
            if dashes and dashes[i]:
                ln = ser.format.line._get_or_add_ln()
                pd = etree.SubElement(ln, qn("a:prstDash"))
                pd.set("val", "dash")
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = 7
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = rgb(col)
            ser.marker.format.line.color.rgb = rgb(col)
            if labels or label_last:
                if label_last:
                    pt = ser.points[len(cats) - 1]
                    dl = pt.data_label
                    dl.has_text_frame = False
                    dl.position = XL_LABEL_POSITION.RIGHT
                    dl.font.size = Pt(size)
                    dl.font.bold = True
                    dl.font.color.rgb = rgb(col)
                    dLbl = dl._dLbl
                    nf = dLbl.find(qn("c:numFmt"))
                    self._dlbl_show(dLbl, fmt)
                else:
                    ser.data_labels.number_format = fmt
                    ser.data_labels.number_format_is_linked = False
                    ser.data_labels.show_value = True
                    ser.data_labels.position = XL_LABEL_POSITION.ABOVE
                    ser.data_labels.font.size = Pt(size - 1)
                    ser.data_labels.font.color.rgb = rgb(col)
                    ser.data_labels.font.bold = True
        return ch

    def _dlbl_show(self, dLbl, fmt):
        # ensure value is shown with the given number format on a single point label
        for tag, val in (("c:showLegendKey", "0"), ("c:showVal", "1"), ("c:showCatName", "0"),
                         ("c:showSerName", "0"), ("c:showPercent", "0"), ("c:showBubbleSize", "0")):
            el = dLbl.find(qn(tag))
            if el is None:
                el = etree.SubElement(dLbl, qn(tag))
            el.set("val", val)
        nf = dLbl.find(qn("c:numFmt"))
        if nf is None:
            nf = etree.Element(qn("c:numFmt"))
            idx = dLbl.find(qn("c:idx"))
            idx.addnext(nf)
        nf.set("formatCode", fmt)
        nf.set("sourceLinked", "0")

    def scatter(self, s, x, y, w, h, points, xr=None, yr=None, xtitle=None, ytitle=None, size=11, title=None):
        """points: list of (name, x, y, color, marker_size)."""
        cd = XyChartData()
        for pt in points:
            name, px, py = pt[0], pt[1], pt[2]
            ser = cd.add_series(name)
            ser.add_data_point(px, py)
        gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        self._style_chart(ch, legend=False, size=size, title=title)
        va, ca = ch.value_axis, ch.category_axis
        for ax, rng, ttl in ((ca, xr, xtitle), (va, yr, ytitle)):
            if rng:
                ax.minimum_scale, ax.maximum_scale = rng[0], rng[1]
                if len(rng) > 2:
                    ax.major_unit = rng[2]
            ax.has_major_gridlines = True
            ax.major_gridlines.format.line.color.rgb = rgb("E3E8EF")
            ax.format.line.color.rgb = rgb("B8C2CE")
            ax.tick_labels.font.size = Pt(10.5)
            ax.tick_labels.font.color.rgb = rgb("muted")
            ax.tick_labels.number_format = "0.0"
            ax.tick_labels.number_format_is_linked = False
            if ttl:
                ax.has_title = True
                ax.axis_title.text_frame.text = ttl
                r0 = ax.axis_title.text_frame.paragraphs[0].runs[0]
                r0.font.size = Pt(11)
                r0.font.bold = True
                r0.font.color.rgb = rgb("navy")
        for i, ser in enumerate(ch.plots[0].series):
            name, px, py, col, ms = points[i][:5]
            pos = points[i][5] if len(points[i]) > 5 else "r"
            ser.format.line.fill.background()
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = ms
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = rgb(col)
            ser.marker.format.line.color.rgb = rgb("white")
            if not name or name.startswith("_"):
                continue
            dl = ser.points[0].data_label
            dl.has_text_frame = True
            dl.text_frame.text = name
            dl.position = {"r": XL_LABEL_POSITION.RIGHT, "l": XL_LABEL_POSITION.LEFT,
                           "t": XL_LABEL_POSITION.ABOVE, "b": XL_LABEL_POSITION.BELOW}[pos]
            rr = dl.text_frame.paragraphs[0].runs[0]
            rr.font.size = Pt(11)
            rr.font.bold = True
            rr.font.color.rgb = rgb(col)
        return ch
